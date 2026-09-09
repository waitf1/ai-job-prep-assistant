from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document


class DocumentParseError(RuntimeError):
    """Raised when an uploaded document cannot be parsed."""


@dataclass(frozen=True)
class ParsedDocument:
    file_name: str
    file_type: str
    text: str


def parse_document(file_name: str, file_bytes: bytes) -> ParsedDocument:
    """Parse an uploaded PDF, docx, txt, or markdown file into plain text."""

    suffix = Path(file_name).suffix.lower()
    if suffix == ".pdf":
        text = _parse_pdf(file_bytes)
    elif suffix == ".docx":
        text = _parse_docx(file_bytes)
    elif suffix in {".txt", ".md", ".markdown"}:
        text = _parse_text(file_bytes)
    else:
        raise DocumentParseError(f"Unsupported file type: {suffix or 'unknown'}")

    text = _normalize_text(text)
    if not text:
        raise DocumentParseError("No readable text was extracted from the file.")

    return ParsedDocument(file_name=file_name, file_type=suffix.lstrip("."), text=text)


def _parse_pdf(file_bytes: bytes) -> str:
    try:
        with pymupdf.open(stream=file_bytes, filetype="pdf") as document:
            pages = []
            for index, page in enumerate(document, start=1):
                page_text = page.get_text().strip()
                if page_text:
                    pages.append(f"--- Page {index} ---\n{page_text}")
            return "\n\n".join(pages)
    except Exception as exc:
        raise DocumentParseError(f"Failed to parse PDF: {exc}") from exc


def _parse_docx(file_bytes: bytes) -> str:
    try:
        document = Document(BytesIO(file_bytes))
        parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))

        return "\n".join(parts)
    except Exception as exc:
        raise DocumentParseError(f"Failed to parse docx: {exc}") from exc


def _parse_text(file_bytes: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gbk"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("Failed to decode text file with utf-8 or gbk.")


def _normalize_text(text: str) -> str:
    lines = [line.rstrip() for line in text.splitlines()]
    return "\n".join(lines).strip()
